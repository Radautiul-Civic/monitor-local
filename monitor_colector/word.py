"""Textul extras din PDF-uri (strat de text sau OCR), salvat ca documente Word.

Rezultatul e pentru verificare internă (artefactul rularii), nu pentru publicare:
textul poate conține date personale.
"""
import logging
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from . import config

log = logging.getLogger(__name__)

_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def safe_name(s, maxlen=110):
    s = _BAD.sub(" ", s or "").replace("_", " ")
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s[:maxlen].rstrip(" .") or "document"


def _field(paragraph, instr):
    """Câmp Word (PAGE / NUMPAGES), actualizat automat la deschidere."""
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == "1":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        run._r.append(el)


def paragraphs(page_text):
    """Grupează rândurile în paragrafe: rând gol = paragraf nou; cratima de la
    capăt de rând se lipește de cuvântul următor."""
    out, cur = [], []
    for line in page_text.splitlines():
        line = line.strip()
        if not line:
            if cur:
                out.append(cur)
                cur = []
            continue
        cur.append(line)
    if cur:
        out.append(cur)
    res = []
    for lines in out:
        txt = lines[0]
        for ln in lines[1:]:
            txt = txt[:-1] + ln if txt.endswith("-") and not txt.endswith(" -") else txt + " " + ln
        res.append(txt)
    return res


def write_docx(text, dest, title, source_url, ocr):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    doc.add_heading(title, level=1)
    meta = doc.add_paragraph()
    meta.add_run("Sursa: ").bold = True
    meta.add_run(source_url)
    note = doc.add_paragraph()
    note.add_run(
        ("Text obținut automat prin OCR din documentul scanat; poate conține erori de recunoaștere. "
         if ocr else "Text extras automat din stratul de text al PDF-ului. ")
        + "Documentul oficial este PDF-ul de la sursă.").italic = True

    pages = text.split("\f")
    for i, page in enumerate(pages, 1):
        h = doc.add_paragraph()
        h.add_run(f"Pagina {i} din documentul original").bold = True
        for p in paragraphs(page):
            doc.add_paragraph(p)
        if i < len(pages):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Pagina ")
    _field(footer, "PAGE")
    footer.add_run(" din ")
    _field(footer, "NUMPAGES")

    os.makedirs(os.path.dirname(dest), exist_ok=True)
    doc.save(dest)


def _title_for(db, row):
    """Nume lizibil: „HCL 160 din 2026”, „PV ședință ordinară 25.09.2026” etc."""
    h = db.execute("SELECT year, number FROM hcl WHERE pdf_url=?", (row["url"],)).fetchone()
    if h:
        return "hcl", f"HCL {h['number']} din {h['year']}"
    p = db.execute("SELECT source, title FROM posts WHERE wp_id=?", (row["wp_id"],)).fetchone() \
        if row["wp_id"] else None
    if p:
        return p["source"], p["title"] or row["link_text"]
    base = os.path.splitext(os.path.basename(row["url"]))[0]
    return row["origin"] or "altele", row["link_text"] or base


def export_all(store, data_dir=config.DATA_DIR, only_new=True):
    """Creează data/texte/<sursa>/<titlu>.docx pentru fiecare document cu text."""
    db = store.db
    stats = {"word": 0, "existente": 0, "erori": 0}
    used = set()
    for row in db.execute("SELECT url, wp_id, origin, link_text, text_path, text_layer FROM documents "
                          "WHERE text_path IS NOT NULL ORDER BY rowid"):
        src, title = _title_for(db, row)
        name = safe_name(title)
        dest = os.path.join(data_dir, "texte", safe_name(src, 40), name + ".docx")
        k = 2
        while dest in used:
            dest = os.path.join(data_dir, "texte", safe_name(src, 40), f"{name} ({k}).docx")
            k += 1
        used.add(dest)
        if only_new and os.path.exists(dest) and os.path.getmtime(dest) >= os.path.getmtime(row["text_path"]):
            stats["existente"] += 1
            continue
        try:
            with open(row["text_path"], encoding="utf-8") as f:
                text = f.read()
            write_docx(text, dest, title, row["url"], ocr=row["text_layer"] == 0)
            stats["word"] += 1
        except Exception as e:
            log.warning("Word %s: %s", row["url"], e)
            stats["erori"] += 1
    return stats
